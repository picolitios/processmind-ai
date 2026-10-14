import csv
import re
import zipfile
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.core.errors import ApiError

MAX_ROWS = 50_000
MAX_COLUMNS = 64
MAX_CELL = 4_000
MAX_UNCOMPRESSED = 100 * 1024 * 1024
REQUIRED = (
    "purchase_id",
    "supplier_id",
    "purchase_date",
    "product_id",
    "quantity",
    "unit_price",
    "currency",
    "total_amount",
)
OPTIONAL = (
    "supplier_name",
    "product_name",
    "unit_of_measure",
    "expected_delivery_date",
    "actual_delivery_date",
)
FIELDS = (*REQUIRED, *OPTIONAL)


def check_file(path: Path, file_format: str) -> None:
    if file_format == "csv":
        with path.open("rb") as source:
            head = source.read(4096)
        if b"\x00" in head or head.startswith((b"PK\x03\x04", b"\xd0\xcf\x11\xe0")):
            raise ApiError(422, "invalid_file", "CSV inválido")
        return
    if file_format != "xlsx" or not zipfile.is_zipfile(path):
        raise ApiError(422, "invalid_file", "XLSX inválido")
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
            raise ApiError(422, "invalid_file", "XLSX inválido")
        if any("vbaProject" in name or "externalLinks" in name for name in names):
            raise ApiError(422, "unsafe_file", "XLSX contém conteúdo não permitido")
        if (
            len(names) > 1000
            or sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED
        ):
            raise ApiError(413, "file_too_large", "XLSX excede limite de descompressão")
        for item in archive.infolist():
            if item.file_size > 0 and item.file_size > max(1, item.compress_size) * 1000:
                raise ApiError(413, "file_too_large", "XLSX excede limite de compactação")


def csv_options(path: Path) -> tuple[str, str]:
    with path.open("rb") as source:
        sample = source.read(16_384)
    try:
        decoded = sample.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError as exc:
        raise ApiError(422, "invalid_encoding", "CSV deve usar UTF-8") from exc
    try:
        delimiter = csv.Sniffer().sniff(decoded, delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    return encoding, delimiter


def table_rows(
    path: Path, file_format: str, encoding: str | None = None, delimiter: str | None = None
) -> Iterator[list[str]]:
    if file_format == "csv":
        detected_encoding, detected_delimiter = csv_options(path)
        if encoding not in (None, "utf-8", "utf-8-sig"):
            raise ApiError(422, "invalid_encoding", "CSV deve usar UTF-8")
        if delimiter not in (None, ",", ";", "\t"):
            raise ApiError(422, "invalid_delimiter", "Delimitador inválido")
        try:
            with path.open("r", encoding=encoding or detected_encoding, newline="") as source:
                reader = csv.reader(source, delimiter=delimiter or detected_delimiter, strict=True)
                yield from reader
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ApiError(422, "invalid_file", "CSV inválido") from exc
    else:
        try:
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet is None:
                    raise ApiError(422, "invalid_file", "Planilha vazia")
                for row in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise ApiError(422, "unsafe_file", "Fórmulas não são permitidas")
                    yield ["" if cell.value is None else str(cell.value) for cell in row]
            finally:
                workbook.close()
        except (ValueError, KeyError, zipfile.BadZipFile) as exc:
            raise ApiError(422, "invalid_file", "XLSX inválido") from exc


def bounded_rows(rows: Iterator[list[str]]) -> Iterator[tuple[int, list[str]]]:
    for number, row in enumerate(rows, start=1):
        if number > MAX_ROWS + 1:
            raise ApiError(413, "too_many_rows", "Limite de 50 mil linhas excedido")
        if len(row) > MAX_COLUMNS or any(len(value) > MAX_CELL for value in row):
            raise ApiError(413, "table_limit", "Limite de colunas ou células excedido")
        yield number, row


def preview(path: Path, file_format: str) -> dict[str, Any]:
    check_file(path, file_format)
    encoding, delimiter = csv_options(path) if file_format == "csv" else (None, None)
    rows = bounded_rows(table_rows(path, file_format, encoding, delimiter))
    try:
        _, headers = next(rows)
    except StopIteration as exc:
        raise ApiError(422, "empty_file", "Arquivo vazio") from exc
    if (
        not headers
        or len(headers) != len(set(headers))
        or any(not item.strip() for item in headers)
    ):
        raise ApiError(422, "invalid_headers", "Cabeçalhos vazios ou repetidos")
    sample: list[list[str]] = []
    for _, row in rows:
        if len(row) != len(headers):
            raise ApiError(422, "invalid_row", "Número de colunas inconsistente")
        if len(sample) < 20:
            sample.append(row)
        if len(sample) == 20:
            break
    return {
        "headers": headers,
        "rows": sample,
        "detected_encoding": encoding,
        "delimiter": delimiter,
    }


def validate_mapping(columns: dict[str, str], headers: list[str]) -> None:
    if not set(REQUIRED) <= columns.keys() or not columns.keys() <= set(FIELDS):
        raise ApiError(422, "invalid_mapping", "Campos obrigatórios ausentes ou desconhecidos")
    if len(set(columns.values())) != len(columns) or not set(columns.values()) <= set(headers):
        raise ApiError(422, "invalid_mapping", "Colunas de origem inválidas ou repetidas")


def normalize(
    raw: dict[str, str], mapping: dict[str, str]
) -> tuple[dict[str, str], list[dict[str, str]]]:
    values = {field: raw.get(source, "").strip() for field, source in mapping.items()}
    errors: list[dict[str, str]] = []

    def error(field: str, code: str) -> None:
        errors.append({"field": field, "code": code, "message": f"Valor inválido para {field}"})

    for field in ("purchase_id", "supplier_id", "product_id"):
        if not values.get(field) or len(values[field]) > 160:
            error(field, "required_or_too_long")
    for field in ("supplier_name", "product_name"):
        if len(values.get(field, "")) > 255:
            error(field, "too_long")
    if len(values.get("unit_of_measure", "")) > 40:
        error("unit_of_measure", "too_long")
    for field in ("purchase_date", "expected_delivery_date", "actual_delivery_date"):
        value = values.get(field)
        if not value and field == "purchase_date":
            error(field, "required")
        elif value:
            try:
                values[field] = date.fromisoformat(value[:10]).isoformat()
                if (
                    len(value) > 10
                    and datetime.fromisoformat(value).date().isoformat() != values[field]
                ):
                    error(field, "invalid_date")
            except ValueError:
                error(field, "invalid_date")
    for field, scale in (("quantity", 4), ("unit_price", 4), ("total_amount", 2)):
        try:
            decimal_value = Decimal(values.get(field, ""))
            exponent = decimal_value.as_tuple().exponent
            max_integer_digits = 18 if field == "total_amount" else 14
            if (
                not decimal_value.is_finite()
                or (isinstance(exponent, int) and exponent < -scale)
                or decimal_value.adjusted() >= max_integer_digits
            ):
                raise InvalidOperation
            if decimal_value <= 0:
                error(field, "non_positive")
            values[field] = str(decimal_value)
        except (InvalidOperation, ValueError):
            error(field, "invalid_decimal")
    currency = values.get("currency", "").upper()
    if not re.fullmatch(r"[A-Z]{3}", currency):
        error("currency", "invalid_currency")
    values["currency"] = currency
    return values, errors
